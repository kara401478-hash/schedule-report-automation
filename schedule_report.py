"""
スケジュールデータ集計・レポート自動生成スクリプト
社内システムから複数種別のスケジュールデータをCSV取得し、
担当者別の件数集計Excelを自動生成。OneDriveへのコピーとメール送信まで自動化。
タスクスケジューラでの定期実行を想定。
"""

import os
import csv
import time
import shutil
import traceback
from datetime import datetime, timedelta
from dotenv import load_dotenv

import openpyxl
import pandas as pd
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager

# ============================================================
# 設定読み込み
# ============================================================
load_dotenv()

LOGIN_ID       = os.getenv("LOGIN_ID")
LOGIN_PASSWORD = os.getenv("LOGIN_PASSWORD")
TARGET_URL     = os.getenv("TARGET_URL")
NAV_LABEL      = os.getenv("NAV_LABEL")
DOWNLOAD_FOLDER = os.getenv("DOWNLOAD_FOLDER")
WORK_EXCEL_PATH  = os.getenv("WORK_EXCEL_PATH")
OUTPUT_EXCEL_PATH = os.getenv("OUTPUT_EXCEL_PATH")
ONEDRIVE_PATH  = os.getenv("ONEDRIVE_PATH")
MAIL_TO        = os.getenv("MAIL_TO")
MAIL_CC        = os.getenv("MAIL_CC")
MAIL_SUBJECT   = os.getenv("MAIL_SUBJECT")
ONEDRIVE_URL   = os.getenv("ONEDRIVE_URL")

# スケジュール種別インデックス（システムのドロップダウン順に合わせて設定）
SCHEDULE_TYPE_INDEX_1 = 14   # スケジュール種別1に対応するオプションのインデックス
SCHEDULE_TYPE_INDEX_2 = 18   # スケジュール種別2に対応するオプションのインデックス

# 所属の並び順（レポートの表示順）
SORT_ORDER = [
    "branch_A",
    "branch_B",
    "branch_C",
    # 必要に応じて追加
]

# ============================================================
# ユーティリティ関数
# ============================================================

def setup_driver(download_folder):
    """ChromeDriverのセットアップ"""
    options = Options()
    prefs = {
        "download.default_directory": download_folder,
        "download.prompt_for_download": False,
        "directory_upgrade": True,
        "safebrowsing.enabled": True,
    }
    options.add_experimental_option("prefs", prefs)
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)
    return driver


def wait_for_downloads(download_folder, timeout=60):
    """ダウンロード完了まで待機"""
    for _ in range(timeout):
        time.sleep(1)
        if not any(f.endswith(".crdownload") for f in os.listdir(download_folder)):
            print("Download complete.")
            return True
    print("Timeout: Download did not complete.")
    return False


def rename_latest_csv(download_folder, new_file_name):
    """最新のCSVを指定ファイル名にリネーム"""
    new_path = os.path.join(download_folder, new_file_name)
    csv_files = [f for f in os.listdir(download_folder)
                 if f.endswith(".csv") and f != new_file_name]
    if not csv_files:
        print(f"CSV not found for: {new_file_name}")
        return False
    latest = max(
        [os.path.join(download_folder, f) for f in csv_files],
        key=os.path.getctime
    )
    if os.path.exists(new_path):
        os.remove(new_path)
    os.rename(latest, new_path)
    print(f"Saved: {new_path}")
    return True


def download_csv_from_menu(wait, driver, download_folder, new_file_name, wait_time=30):
    """CSVダウンロードメニューから共通ダウンロード処理"""
    time.sleep(wait_time)
    try:
        WebDriverWait(driver, 20).until(
            EC.element_to_be_clickable((By.XPATH, '//span[text()="CSVダウンロード"]'))
        ).click()
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, '/html/body/div[2]/div/div[2]/div[1]/div/div[2]/div/a[2]')
        )).click()
        if wait_for_downloads(download_folder):
            return rename_latest_csv(download_folder, new_file_name)
    except Exception as e:
        print(f"Error during CSV download: {e}")
    return False


# ============================================================
# Step1: CSVダウンロード
# ============================================================

def download_schedule_data(driver, wait):
    """スケジュールデータ（種別1・種別2・人員）をCSVでダウンロード"""
    today = datetime.today()
    first_day_str = today.replace(day=1).strftime("%Y/%m/%d")
    today_str = today.strftime("%Y/%m/%d")

    # ログイン
    driver.get(TARGET_URL)
    wait.until(EC.presence_of_element_located(
        (By.XPATH, '/html/body/div/form/div[1]/input')
    )).send_keys(LOGIN_ID)
    driver.find_element(By.XPATH, '/html/body/div/form/div[2]/input').send_keys(LOGIN_PASSWORD)
    wait.until(EC.element_to_be_clickable(
        (By.XPATH, '/html/body/div/form/input[2]')
    )).click()
    print("✅ ログイン完了")

    for type_index, file_name in [
        (SCHEDULE_TYPE_INDEX_1, "1.schedule_type1.csv"),
        (SCHEDULE_TYPE_INDEX_2, "2.schedule_type2.csv"),
    ]:
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, f"//span[@class='nav-label' and text()='{NAV_LABEL}']")
        )).click()
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//button[contains(text(),'条件を指定して検索')]")
        )).click()

        wait.until(EC.element_to_be_clickable((By.ID, "q_start_date_gteq"))).clear()
        wait.until(EC.element_to_be_clickable((By.ID, "q_start_date_gteq"))).send_keys(first_day_str)
        wait.until(EC.element_to_be_clickable((By.ID, "q_end_date_lteq"))).clear()
        wait.until(EC.element_to_be_clickable((By.ID, "q_end_date_lteq"))).send_keys(today_str)

        select = wait.until(EC.element_to_be_clickable((By.ID, "q_event_kind_id_eq")))
        select.find_elements(By.TAG_NAME, "option")[type_index].click()

        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "/html/body/div[2]/div/div[2]/div[2]/div/div/form/div[3]/input")
        )).click()
        time.sleep(15)

        btn = wait.until(EC.element_to_be_clickable(
            (By.XPATH, '/html/body/div[2]/div/div[2]/div[1]/div/div[2]/div/a[2]')
        ))
        driver.execute_script("arguments[0].scrollIntoView(true);", btn)
        driver.execute_script("arguments[0].click();", btn)
        download_csv_from_menu(wait, driver, DOWNLOAD_FOLDER, file_name)
        print(f"✅ {file_name} ダウンロード完了")

    # 人員データ
    driver.find_element(By.LINK_TEXT, "設定").click()
    WebDriverWait(driver, 30).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, 'a[href="/admins/imports/new"]'))
    ).click()
    time.sleep(5)
    WebDriverWait(driver, 30).until(
        EC.element_to_be_clickable((By.XPATH, '//a[text()="人員"]'))
    ).click()
    time.sleep(5)

    try:
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//div[contains(text(),'CSVダウンロード')]")
        )).click()
    except Exception:
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, '/html/body/div[2]/div/div[3]/div/div[2]/div/div[2]/div/div[1]/div')
        )).click()

    time.sleep(3)
    try:
        btn = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//a[contains(text(),'CSVダウンロード')]")
        ))
        driver.execute_script("arguments[0].click();", btn)
    except Exception:
        wait.until(EC.element_to_be_clickable(
            (By.XPATH, '/html/body/div[2]/div/div[3]/div/div[2]/div/div[2]/div/div[1]/div/ul/li[5]/a')
        )).click()

    download_csv_from_menu(wait, driver, DOWNLOAD_FOLDER, "3.staff_list.csv")
    print("✅ 人員データ ダウンロード完了")


# ============================================================
# Step2: データ加工・Excel生成
# ============================================================

def add_source_column(file_path, source_name):
    """CSVにSource.Name列を追加"""
    df = pd.read_csv(file_path, encoding="cp932")
    if "Source.Name" not in df.columns:
        df.insert(0, "Source.Name", source_name)
    df.to_csv(file_path, index=False, encoding="cp932")


def dedup_staff_list(csv_path, output_path, name_column="名前::name"):
    """スタッフリストの重複を削除してExcel保存"""
    df = pd.read_csv(csv_path, encoding="cp932")
    if name_column in df.columns:
        df = df.drop_duplicates(subset=[name_column], keep="last")
    df.to_excel(output_path, index=False)
    print(f"✅ スタッフリスト重複削除完了: {output_path}")


def build_work_excel(work_excel_path, download_folder, staff_dedup_path):
    """作業用Excelを生成（Sheet1にCSVデータ統合、スタッフリストシート追加、集計シート追加）"""
    csv1 = os.path.join(download_folder, "1.schedule_type1.csv")
    csv2 = os.path.join(download_folder, "2.schedule_type2.csv")

    # Source.Name列を追加
    add_source_column(csv1, "schedule_type1.csv")
    add_source_column(csv2, "schedule_type2.csv")

    # 作業用Excel作成
    with pd.ExcelWriter(work_excel_path, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Sheet1", index=False)
    print(f"✅ 作業用Excel作成: {work_excel_path}")

    wb = openpyxl.load_workbook(work_excel_path)

    # Sheet1にCSVデータを統合
    sheet1 = wb["Sheet1"]
    for cell in sheet1.iter_rows():
        for c in cell:
            c.value = None

    current_row = 1
    for i, csv_path in enumerate([csv1, csv2]):
        with open(csv_path, "r", encoding="cp932") as f:
            reader = csv.reader(f)
            for j, row in enumerate(reader):
                if j == 0 and i > 0:
                    continue
                for col_idx, value in enumerate(row, start=1):
                    sheet1.cell(row=current_row, column=col_idx, value=value)
                current_row += 1

    # スタッフリストシートを追加
    wb_staff = openpyxl.load_workbook(staff_dedup_path)
    staff_sheet_name = "staff_list_dedup"
    if staff_sheet_name not in wb.sheetnames:
        wb.create_sheet(staff_sheet_name)
    new_sheet = wb[staff_sheet_name]
    src = wb_staff.active
    for r in range(1, src.max_row + 1):
        for c in range(1, src.max_column + 1):
            new_sheet.cell(row=r, column=c, value=src.cell(row=r, column=c).value)

    # 集計シートを追加
    summary_name = "summary"
    if summary_name not in wb.sheetnames:
        wb.create_sheet(summary_name)
    summary = wb[summary_name]

    previous_date = datetime.now() - timedelta(days=1)
    summary["G1"] = previous_date
    summary["G1"].number_format = 'yyyy/mm/dd"現在"'

    headers = ["組織", "部門", "所属", "役職", "担当者", "種別1件数", "種別2件数"]
    for col_idx, header in enumerate(headers, start=1):
        summary.cell(row=2, column=col_idx, value=header)

    unique_creators = []
    for row in sheet1.iter_rows(min_row=2, max_row=sheet1.max_row, min_col=19, max_col=19):
        for cell in row:
            if cell.value and cell.value not in unique_creators:
                unique_creators.append(cell.value)

    for i, creator in enumerate(unique_creators, start=3):
        summary[f"E{i}"] = creator
        summary[f"A{i}"] = f'=IFERROR(IF($E{i}="","",INDEX({staff_sheet_name}!J:J,MATCH(summary!$E{i},{staff_sheet_name}!$B:$B,0))),"")'
        summary[f"B{i}"] = f'=IFERROR(IF($E{i}="","",INDEX({staff_sheet_name}!K:K,MATCH(summary!$E{i},{staff_sheet_name}!$B:$B,0))),"")'
        summary[f"C{i}"] = f'=IFERROR(IF($E{i}="","",INDEX({staff_sheet_name}!L:L,MATCH(summary!$E{i},{staff_sheet_name}!$B:$B,0))),"")'
        summary[f"D{i}"] = f'=IFERROR(IF($E{i}="","",INDEX({staff_sheet_name}!M:M,MATCH(summary!$E{i},{staff_sheet_name}!$B:$B,0))),"")'
        summary[f"F{i}"] = f'=IF($E{i}=" "," ",COUNTIFS(Sheet1!$C:$C,summary!$F$2,Sheet1!$S:$S,summary!E{i}))'
        summary[f"G{i}"] = f'=IF($E{i}=" "," ",COUNTIFS(Sheet1!$C:$C,summary!$G$2,Sheet1!$S:$S,summary!E{i}))'

    for col in ["A", "B", "C", "D", "E", "F", "G"]:
        summary.column_dimensions[col].width = 15

    wb.save(work_excel_path)
    print(f"✅ 作業用Excel保存完了: {work_excel_path}")


def build_output_excel(work_excel_path, output_excel_path):
    """集計シートを並び替えて最終Excelとして出力"""
    import xlwings as xw

    app = xw.App(visible=False)
    try:
        wb = app.books.open(work_excel_path)
        source_sheet = wb.sheets["summary"]

        new_wb = xw.Book()
        new_sheet = new_wb.sheets[0]
        new_sheet.name = "summary"

        data = source_sheet.used_range.value
        header_rows = data[0:2]
        data_rows = data[2:]

        def sort_key(row):
            affiliation = row[2] if row and len(row) > 2 else None
            if affiliation in SORT_ORDER:
                return SORT_ORDER.index(affiliation)
            return len(SORT_ORDER)

        sorted_rows = sorted(data_rows, key=sort_key)
        new_sheet.range("A1").value = header_rows + sorted_rows

        previous_date = datetime.now() - timedelta(days=1)
        new_sheet.range("G1").value = previous_date
        new_sheet.range("G1").number_format = 'yyyy/mm/dd"現在"'
        new_sheet.range("A:G").column_width = 15

        last_row = new_sheet.used_range.last_cell.row
        new_sheet.range(f"A3:G{last_row}").api.Borders.LineStyle = 1
        new_sheet.range("G1").api.Borders.LineStyle = 1
        new_sheet.range("A2:G2").api.AutoFilter(1)

        wb.close()
        new_wb.save(output_excel_path)
        print(f"✅ 出力Excel保存完了: {output_excel_path}")
    except Exception as e:
        print(f"❌ 出力Excel生成エラー: {e}")
        traceback.print_exc()
    finally:
        new_wb.close()
        app.quit()


# ============================================================
# Step3: OneDriveコピー & メール送信
# ============================================================

def copy_to_onedrive(source, destination):
    """OneDriveにファイルをコピー"""
    if os.path.exists(source):
        shutil.copy2(source, destination)
        print(f"✅ OneDriveコピー完了: {destination}")
    else:
        raise FileNotFoundError(f"ソースファイルが見つかりません: {source}")


def send_mail():
    """Outlookでメール送信"""
    import win32com.client as win32
    outlook = win32.Dispatch("outlook.application")
    mail = outlook.CreateItem(0)
    mail.To = MAIL_TO
    mail.CC = MAIL_CC
    mail.Subject = MAIL_SUBJECT
    html_body = f"""<html><body>
<p>各位</p>
<p>{MAIL_SUBJECT}の集計をお送りします。</p>
<p><a href="{ONEDRIVE_URL}">レポートを開く</a></p>
<p>以上、ご確認ください。</p>
</body></html>"""
    mail.HTMLBody = html_body
    mail.Send()
    print(f"✅ メール送信完了: {datetime.now()}")


# ============================================================
# メイン処理
# ============================================================

def main():
    driver = setup_driver(DOWNLOAD_FOLDER)
    wait = WebDriverWait(driver, 20)

    try:
        # Step1: CSVダウンロード
        download_schedule_data(driver, wait)
    except Exception as e:
        print(f"❌ ダウンロードエラー: {e}")
        traceback.print_exc()
        driver.save_screenshot(os.path.join(DOWNLOAD_FOLDER, "error.png"))
    finally:
        driver.quit()

    # Step2: データ加工・Excel生成
    staff_csv = os.path.join(DOWNLOAD_FOLDER, "3.staff_list.csv")
    staff_dedup_path = os.path.join(DOWNLOAD_FOLDER, "staff_list_dedup.xlsx")
    dedup_staff_list(staff_csv, staff_dedup_path)
    build_work_excel(WORK_EXCEL_PATH, DOWNLOAD_FOLDER, staff_dedup_path)
    build_output_excel(WORK_EXCEL_PATH, OUTPUT_EXCEL_PATH)

    # Step3: OneDriveコピー & メール送信
    try:
        copy_to_onedrive(OUTPUT_EXCEL_PATH, ONEDRIVE_PATH)
        send_mail()
    except Exception as e:
        print(f"❌ コピー/メール送信エラー: {e}")
        traceback.print_exc()


if __name__ == "__main__":
    main()
